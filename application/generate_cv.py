#!/usr/bin/env python3
"""Render the reviewed Markdown CV to accessible-text PDF and editable DOCX.

Dependencies are confined to document generation; no application secrets or
environment variables are read. Candidate contact information is read only
from the reviewed cv-content.md source. Run from any directory. This is intentionally
a small Markdown subset: headings, paragraphs, and single-level bullet lists.
"""

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate


DIRECTORY = Path(__file__).resolve().parent
BASENAME = "Eldar_Smagulov_SerpApi_Junior_Fullstack_CV"


def read_blocks():
    """Preserve deliberate separate source lines without interpreting HTML."""
    for raw_line in (DIRECTORY / "cv-content.md").read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            yield "title", line[2:]
        elif line.startswith("## "):
            yield "section", line[3:]
        elif line.startswith("- "):
            yield "bullet", line[2:]
        else:
            yield "paragraph", line


def pdf_fonts():
    roots = [Path("/usr/share/fonts/truetype/dejavu"), Path("/Library/Fonts")]
    for root in roots:
        regular, bold = root / "DejaVuSans.ttf", root / "DejaVuSans-Bold.ttf"
        if regular.is_file() and bold.is_file():
            pdfmetrics.registerFont(TTFont("CVSans", str(regular)))
            pdfmetrics.registerFont(TTFont("CVSansBold", str(bold)))
            return "CVSans", "CVSansBold"
    return "Helvetica", "Helvetica-Bold"


def create_pdf(blocks):
    regular, bold = pdf_fonts()
    common = {"fontName": regular, "fontSize": 9.1, "leading": 12.0, "alignment": TA_LEFT}
    styles = {
        "title": ParagraphStyle("CVTitle", fontName=bold, fontSize=21, leading=25, spaceAfter=5),
        "section": ParagraphStyle("CVSection", fontName=bold, fontSize=10, leading=13, spaceBefore=9, spaceAfter=4, textColor=colors.HexColor("#17304c")),
        "paragraph": ParagraphStyle("CVParagraph", **common, spaceAfter=3),
        "bullet": ParagraphStyle("CVBullet", **common, leftIndent=10, firstLineIndent=-7, spaceAfter=3),
    }
    story = []
    for kind, content in blocks:
        safe = escape(content)
        if kind == "bullet":
            safe = "&#8226; " + safe
        story.append(Paragraph(safe, styles[kind]))
    document = SimpleDocTemplate(
        str(DIRECTORY / f"{BASENAME}.pdf"), pagesize=A4,
        rightMargin=45, leftMargin=45, topMargin=35, bottomMargin=35,
        title="Eldar Smagulov — Junior Fullstack Engineer CV",
        author="Eldar Smagulov", subject="Personal-project application for SerpApi",
    )
    document.build(story)


def create_docx(blocks):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.49)
    section.left_margin = section.right_margin = Inches(0.63)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    document.core_properties.author = "Eldar Smagulov"
    document.core_properties.last_modified_by = "Eldar Smagulov"
    document.core_properties.title = "Eldar Smagulov — Junior Fullstack Engineer CV"
    document.core_properties.subject = "Personal-project application for SerpApi"
    for kind, content in blocks:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(content if kind != "bullet" else "• " + content)
        if kind == "title":
            run.bold, run.font.size = True, Pt(21)
            paragraph.paragraph_format.space_after = Pt(5)
        elif kind == "section":
            run.bold, run.font.size = True, Pt(10)
            paragraph.paragraph_format.space_before = Pt(9)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.keep_with_next = True
        elif kind == "bullet":
            paragraph.paragraph_format.left_indent = Inches(0.14)
            paragraph.paragraph_format.first_line_indent = Inches(-0.10)
        elif " — " in content:
            run.bold = True
            paragraph.paragraph_format.keep_with_next = True
    document.save(DIRECTORY / f"{BASENAME}.docx")


def main():
    blocks = list(read_blocks())
    create_pdf(blocks)
    create_docx(blocks)
    print(f"Generated {BASENAME}.pdf and {BASENAME}.docx")


if __name__ == "__main__":
    main()
